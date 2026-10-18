"""Editable PowerPoint version of a rendered slide deck (docs/slides/<deck>/).

The HTML stays the single source of the wording: Chrome lays the deck out in
print mode at 1440 x 810 px with Arial forced (the font the PowerPoint uses, so
lines break in the same places), and EXTRACT_JS reads each slide back as
blocks — text boxes, decorated boxes with their text, tables and images —
with their positions. build() turns those into native PowerPoint objects: one
text box per paragraph block, real tables, shapes with text inside, pictures.
"""

import json
import re
from pathlib import Path
from urllib.parse import unquote, urlparse

from lxml import etree
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

VIEW_W, VIEW_H = 1440, 810
SLIDE_W, SLIDE_H = 12192000, 6858000  # 13.333 x 7.5 in
EMU = SLIDE_W / VIEW_W  # EMU per CSS px
PT = 960 / VIEW_W  # pt per CSS px
FONT = "Arial"
MONO = "Courier New"
NO_STYLE_TABLE = (
    "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"
)

FORCE_FONT_CSS = f"""
body, body * {{ font-family: {FONT}, sans-serif !important; }}
code, code * {{ font-family: "{MONO}", monospace !important; }}
"""

EXTRACT_JS = r"""
() => {
const PX = v => parseFloat(v) || 0;
const parse = c => {
  const m = (c || "").match(/rgba?\(([^)]+)\)/);
  if (!m) return null;
  const p = m[1].split(",").map(s => parseFloat(s));
  if (p.length === 4 && p[3] < 0.05) return null;
  return p.slice(0, 3).map(Math.round);
};
const hex = rgb => rgb ? rgb.map(v => v.toString(16).padStart(2, "0")).join("") : null;
const lum = rgb => rgb ? (0.299 * rgb[0] + 0.587 * rgb[1] + 0.114 * rgb[2]) / 255 : 0;
const BLOCK = new Set(["block", "list-item", "flex", "grid", "table", "table-row",
  "table-cell", "flow-root"]);
let origin = null;
const rel = r => ({x: r.left - origin.left, y: r.top - origin.top, w: r.width, h: r.height});

function borders(cs) {
  const out = {};
  for (const s of ["Top", "Right", "Bottom", "Left"]) {
    const w = PX(cs["border" + s + "Width"]);
    const c = parse(cs["border" + s + "Color"]);
    if (w > 0 && c && cs["border" + s + "Style"] !== "none") out[s[0].toLowerCase()] = {w, c: hex(c)};
  }
  return out;
}
function deco(cs) {
  const bg = parse(cs.backgroundColor);
  const b = borders(cs);
  if (!bg && !Object.keys(b).length) return null;
  const radii = ["TopLeft", "TopRight", "BottomRight", "BottomLeft"].map(k => PX(cs["border" + k + "Radius"]));
  return {bg: hex(bg), borders: b, radius: radii.every(r => r === radii[0]) ? radii[0] : 0};
}
const hidden = cs => cs.display === "none" || cs.visibility === "hidden";
const isBlock = cs => BLOCK.has(cs.display);

// can this element's content live inside one text frame?
function flowable(el) {
  for (const d of el.querySelectorAll("*")) {
    const cs = getComputedStyle(d);
    if (hidden(cs)) continue;
    if (["IMG", "TABLE", "SVG", "CANVAS", "VIDEO"].includes(d.tagName)) return false;
    if (["flex", "grid", "inline-flex", "inline-grid"].includes(cs.display)) return false;
    if (cs.position === "absolute" || cs.position === "fixed") return false;
    if (deco(cs)) return false;
  }
  return true;
}
const hasBlockChild = el => [...el.children].some(c => isBlock(getComputedStyle(c)));
const hasText = el => /\S/.test(el.textContent);

function runColor(el, root) {
  let c = parse(getComputedStyle(el).color);
  if (lum(c) > 0.95) {  // white text on a coloured pill inside a frame: use the pill colour
    for (let a = el; a && a !== root; a = a.parentElement) {
      const bg = parse(getComputedStyle(a).backgroundColor);
      if (bg && lum(bg) < 0.9) { c = bg; break; }
    }
  }
  return hex(c);
}
function transform(t, cs) {
  if (cs.textTransform === "uppercase") return t.toUpperCase();
  if (cs.textTransform === "lowercase") return t.toLowerCase();
  return t;
}

// paragraphs (with styled runs) of everything inside root (or just `nodes`)
function paras(root, nodes) {
  const out = [];
  let cur = null;
  const open = (blockEl) => {
    const cs = getComputedStyle(blockEl);
    cur = {align: cs.textAlign, lh: cs.lineHeight === "normal" ? null : PX(cs.lineHeight),
           before: 0, after: 0, bullet: null, indent: 0, runs: []};
    out.push(cur);
    return cur;
  };
  function walk(list, blockEl) {
    for (const n of list) {
      if (n.nodeType === 3) {
        if (!cur) open(blockEl);
        const ps = n.parentElement, cs = getComputedStyle(ps);
        const ls = cs.letterSpacing === "normal" ? 0 : PX(cs.letterSpacing);
        const a = ps.closest("a");
        cur.runs.push({t: transform(n.data, cs), size: PX(cs.fontSize),
          bold: parseInt(cs.fontWeight) >= 600, italic: cs.fontStyle === "italic",
          color: runColor(ps, root), ls, mono: /courier|mono/i.test(cs.fontFamily),
          href: a && /^(https?|mailto):/.test(a.href) ? a.href : null,
          pre: /pre/.test(cs.whiteSpace) && cs.whiteSpace !== "nowrap"});
      } else if (n.nodeType === 1) {
        const cs = getComputedStyle(n);
        if (hidden(cs)) continue;
        if (n.tagName === "BR") { if (!cur) open(blockEl); cur.runs.push({br: true}); continue; }
        if (isBlock(cs)) {
          const p = open(n);
          p.before = PX(cs.marginTop) + PX(cs.paddingTop);
          if (cs.display === "list-item" && cs.listStyleType !== "none") {
            p.bullet = "•";
            p.indent = PX(getComputedStyle(n.parentElement).paddingLeft);
          }
          walk(n.childNodes, n);
          out[out.length - 1].after = PX(cs.marginBottom) + PX(cs.paddingBottom);
          cur = null;
        } else walk(n.childNodes, blockEl);
      }
    }
  }
  open(root);
  walk(nodes || root.childNodes, root);
  // collapse white space like HTML does
  for (const p of out) {
    let prevSpace = true;
    for (const r of p.runs) {
      if (r.br) { prevSpace = true; continue; }
      if (r.pre) { prevSpace = /\s$/.test(r.t); continue; }
      let t = r.t.replace(/\s+/g, " ");
      if (prevSpace) t = t.replace(/^ /, "");
      if (t) prevSpace = t.endsWith(" ");
      r.t = t;
    }
    for (let i = p.runs.length - 1; i >= 0; i--) {  // trim trailing space
      const r = p.runs[i];
      if (r.br) continue;
      r.t = r.t.replace(/ $/, "");
      if (r.t) break;
    }
    p.runs = p.runs.filter(r => r.br || r.t);
    while (p.runs.length && p.runs[p.runs.length - 1].br) p.runs.pop();
  }
  return out.filter(p => p.runs.length);
}

// did the browser wrap any line? (if not, PowerPoint should not wrap either)
function softWrapped(el, ps) {
  const range = document.createRange();
  range.selectNodeContents(el);
  const bottoms = [];
  for (const r of range.getClientRects()) {
    if (r.height < 2 || r.width < 1) continue;
    if (!bottoms.some(b => Math.abs(b - r.bottom) < r.height * 0.5)) bottoms.push(r.bottom);
  }
  const hard = ps.length + ps.reduce((n, p) => n + p.runs.filter(r => r.br).length, 0);
  return bottoms.length > hard;
}
function textBlock(el, cs) {
  const ps = paras(el);
  if (!ps.length) return null;
  const pad = ["Top", "Right", "Bottom", "Left"].map(s => PX(cs["padding" + s]) + PX(cs["border" + s + "Width"]));
  return {paras: ps, wrap: softWrapped(el, ps), pad,
          anchor: cs.display === "table-cell" ? cs.verticalAlign
            : cs.display.startsWith("inline") ? "middle" : "top"};
}

function table(t) {
  const tcs = getComputedStyle(t);
  const tbg = parse(tcs.backgroundColor);
  const rows = [];
  for (const tr of t.rows) {
    const rcs = getComputedStyle(tr);
    if (hidden(rcs)) continue;
    const rbg = parse(rcs.backgroundColor);
    rows.push({rect: rel(tr.getBoundingClientRect()), cells: [...tr.cells].map(td => {
      const cs = getComputedStyle(td);
      const tb = textBlock(td, cs) || {paras: [], wrap: false, pad: [0, 0, 0, 0], anchor: "middle"};
      tb.pad = ["Top", "Right", "Bottom", "Left"].map(s => PX(cs["padding" + s]));
      return {rect: rel(td.getBoundingClientRect()), span: td.colSpan,
              bg: hex(parse(cs.backgroundColor) || rbg || tbg), borders: borders(cs), ...tb};
    })});
  }
  return {kind: "table", rect: rel(t.getBoundingClientRect()), rows, borders: borders(tcs)};
}

function visit(el, out) {
  const cs = getComputedStyle(el);
  if (hidden(cs)) return;
  if (el.tagName === "IMG") {
    out.push({kind: "img", rect: rel(el.getBoundingClientRect()), src: el.currentSrc,
              alt: el.alt, borders: borders(cs)});
    return;
  }
  if (el.tagName === "TABLE") { out.push(table(el)); return; }
  const d = deco(cs);
  const flow = flowable(el);
  const r = rel(el.getBoundingClientRect());
  if (d && flow) { out.push({kind: "box", rect: r, deco: d, text: hasText(el) ? textBlock(el, cs) : null,
                             name: el.className || el.tagName}); return; }
  if (d) out.push({kind: "box", rect: r, deco: d, text: null, name: el.className || el.tagName});
  if (!d && flow && hasText(el) && (!hasBlockChild(el) || ["UL", "OL"].includes(el.tagName))) {
    const tb = textBlock(el, cs);
    const pad = tb.pad;  // text box = content box (lists keep their indent as bullet margin)
    const left = ["UL", "OL"].includes(el.tagName) ? 0 : pad[3];
    out.push({kind: "text", rect: {x: r.x + left, y: r.y + pad[0], w: r.w - pad[1] - left, h: r.h - pad[0] - pad[2]},
              text: {...tb, pad: [0, 0, 0, 0]}, name: el.className || el.tagName.toLowerCase()});
    return;
  }
  // children; runs of loose inline content become their own text boxes
  let group = [];
  const flush = () => {
    if (!group.some(n => /\S/.test(n.textContent))) { group = []; return; }
    const range = document.createRange();
    range.setStartBefore(group[0]); range.setEndAfter(group[group.length - 1]);
    const ps = paras(el, group);
    const gr = rel(range.getBoundingClientRect());
    out.push({kind: "text", rect: gr, text: {paras: ps, wrap: true, pad: [0, 0, 0, 0], anchor: "top"},
              name: "text"});
    group = [];
  };
  for (const n of el.childNodes) {
    if (n.nodeType === 3) { group.push(n); continue; }
    if (n.nodeType !== 1) continue;
    const ccs = getComputedStyle(n);
    if (!isBlock(ccs) && !deco(ccs) && n.tagName !== "IMG" && flowable(n)) { group.push(n); continue; }
    flush();
    visit(n, out);
  }
  flush();
}

const slides = [];
for (const s of document.querySelectorAll("section.slide")) {
  origin = s.getBoundingClientRect();
  const out = [];
  for (const c of s.children) visit(c, out);
  slides.push({label: s.dataset.label, items: out});
}
return slides;
}
"""


def extract(index_html: Path, chrome: str) -> list[dict]:
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=chrome)
        page = browser.new_page(viewport={"width": VIEW_W, "height": VIEW_H})
        page.emulate_media(media="print")
        page.goto(index_html.resolve().as_uri())
        page.add_style_tag(content=FORCE_FONT_CSS)
        page.evaluate("document.fonts.ready")
        slides = page.evaluate(EXTRACT_JS)
        browser.close()
    return slides


# --------------------------------------------------------------------------
# PowerPoint
# --------------------------------------------------------------------------


def _emu(v: float) -> Emu:
    return Emu(round(v * EMU))


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.upper())


ALIGN = {
    "left": PP_ALIGN.LEFT,
    "start": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "end": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}
ANCHOR = {
    "top": MSO_ANCHOR.TOP,
    "middle": MSO_ANCHOR.MIDDLE,
    "bottom": MSO_ANCHOR.BOTTOM,
}


def _bullet(p, char: str, indent_px: float):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(round(indent_px * EMU)))
    pPr.set("indent", str(-round(indent_px * EMU)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", char)


def fill_text(tf, block: dict):
    """Write paragraphs/runs into a text frame."""
    tf.word_wrap = block["wrap"]
    tf.auto_size = MSO_AUTO_SIZE.NONE
    top, right, bottom, left = block["pad"]
    tf.margin_top, tf.margin_right = _emu(top), _emu(right)
    tf.margin_bottom, tf.margin_left = _emu(bottom), _emu(left)
    tf.vertical_anchor = ANCHOR.get(block["anchor"], MSO_ANCHOR.TOP)
    for i, para in enumerate(block["paras"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN.get(para["align"], PP_ALIGN.LEFT)
        if para["lh"]:
            p.line_spacing = Pt(round(para["lh"] * PT * 2) / 2)
        if i and para["before"]:
            p.space_before = Pt(round(para["before"] * PT, 1))
        if i < len(block["paras"]) - 1 and para["after"]:
            p.space_after = Pt(round(para["after"] * PT, 1))
        if para["bullet"]:
            _bullet(p, para["bullet"], para["indent"])
        for run in para["runs"]:
            if run.get("br"):
                p.add_line_break()
                continue
            r_ = p.add_run()
            r_.text = run["t"]
            f = r_.font
            f.name = MONO if run["mono"] else FONT
            f.size = Pt(round(run["size"] * PT * 2) / 2)
            f.bold = run["bold"]
            f.italic = run["italic"]
            if run["color"]:
                f.color.rgb = _rgb(run["color"])
            if run["ls"]:
                r_._r.get_or_add_rPr().set(
                    "spc", str(round(run["ls"] * PT * 100))
                )
            if run["href"]:
                r_.hyperlink.address = run["href"]


def _no_shadow(shape):
    shape.shadow.inherit = False


def _rect(shapes, x, y, w, h, color, radius=0.0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = shapes.add_shape(kind, _emu(x), _emu(y), _emu(w), _emu(h))
    if radius:
        s.adjustments[0] = min(0.5, radius / max(1, min(w, h)))
    if color:
        s.fill.solid()
        s.fill.fore_color.rgb = _rgb(color)
    else:
        s.fill.background()
    s.line.fill.background()
    _no_shadow(s)
    return s


def add_box(shapes, item: dict):
    """A decorated element: fill + outline, odd borders as bars, text inside."""
    x, y, w, h = (item["rect"][k] for k in "xywh")
    d = item["deco"]
    bs = d["borders"]
    # the most common border becomes the outline; the others become bars
    common = None
    if bs:
        keys = [(v["w"], v["c"]) for v in bs.values()]
        best = max(set(keys), key=keys.count)
        if keys.count(best) >= 3 or len(bs) == 4 and len(set(keys)) == 1:
            common = best
    bars = {
        s: v
        for s, v in bs.items()
        if common is None or (v["w"], v["c"]) != common
    }
    group = shapes.add_group_shape() if bars else None
    target = group.shapes if group else shapes
    s = _rect(target, x, y, w, h, d["bg"], d["radius"])
    s.name = item.get("name") or "box"
    if common:
        s.line.color.rgb = _rgb(common[1])
        s.line.width = _emu(common[0])
    for side, v in bars.items():
        bw = v["w"]
        geo = {
            "t": (x, y, w, bw),
            "b": (x, y + h - bw, w, bw),
            "l": (x, y, bw, h),
            "r": (x + w - bw, y, bw, h),
        }[side]
        _rect(target, *geo, v["c"])
    if item.get("text"):
        fill_text(s.text_frame, item["text"])
    return s


def add_text(shapes, item: dict):
    x, y, w, h = (item["rect"][k] for k in "xywh")
    block = item["text"]
    if block[
        "wrap"
    ]:  # a little slack so PowerPoint does not break lines earlier
        align = block["paras"][0]["align"]
        extra = max(4.0, w * 0.03)
        x -= (
            extra / 2
            if align == "center"
            else extra
            if align in ("right", "end")
            else 0
        )
        w += extra
    tb = shapes.add_textbox(_emu(x), _emu(y), _emu(w), _emu(h))
    tb.name = item.get("name") or "text"
    fill_text(tb.text_frame, block)
    return tb


def _cell_border(cell, side: str, b: dict | None):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {"l": "a:lnL", "r": "a:lnR", "t": "a:lnT", "b": "a:lnB"}[side]
    for e in tcPr.findall(qn(tag)):
        tcPr.remove(e)
    ln = etree.Element(qn(tag))
    if b:
        ln.set("w", str(round(b["w"] * EMU)))
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", b["c"].upper())
    else:
        ln.set("w", "0")
        etree.SubElement(ln, qn("a:noFill"))
    # schema order inside tcPr: lnL, lnR, lnT, lnB, ..., fill
    order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
    idx = 0
    for i, child in enumerate(tcPr):
        if child.tag in [qn(t) for t in order[: order.index(tag)]]:
            idx = i + 1
    tcPr.insert(idx, ln)


def add_table(shapes, item: dict):
    rows = item["rows"]
    ncols = max(sum(c["span"] for c in r["cells"]) for r in rows)
    x, y, w, h = (item["rect"][k] for k in "xywh")
    gf = shapes.add_table(len(rows), ncols, _emu(x), _emu(y), _emu(w), _emu(h))
    gf.name = "table"
    tbl = gf.table
    tbl.first_row = tbl.horz_banding = False
    tbl._tbl.tblPr.find(qn("a:tableStyleId")).text = NO_STYLE_TABLE
    # column widths from the row with the most cells (spanned cells split evenly)
    full = max(rows, key=lambda r: len(r["cells"]))
    j = 0
    for c in full["cells"]:
        for k in range(c["span"]):
            tbl.columns[j + k].width = _emu(c["rect"]["w"] / c["span"])
        j += c["span"]
    tb = item["borders"]
    for i, r in enumerate(rows):
        tbl.rows[i].height = _emu(r["rect"]["h"])
        j = 0
        for c in r["cells"]:
            cell = tbl.cell(i, j)
            if c["span"] > 1:
                cell.merge(tbl.cell(i, j + c["span"] - 1))
            if c["bg"]:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(c["bg"])
            else:
                cell.fill.background()
            if c["paras"]:
                fill_text(cell.text_frame, c)
            else:
                cell.text_frame.text = ""
            t, rgt, b, lft = c["pad"]
            cell.margin_top, cell.margin_right = _emu(t), _emu(rgt)
            cell.margin_bottom, cell.margin_left = _emu(b), _emu(lft)
            cell.vertical_anchor = ANCHOR.get(c["anchor"], MSO_ANCHOR.MIDDLE)
            cb = c["borders"]
            last_col = j + c["span"] == ncols
            _cell_border(
                cell, "l", cb.get("l") or (tb.get("l") if j == 0 else None)
            )
            _cell_border(
                cell, "r", cb.get("r") or (tb.get("r") if last_col else None)
            )
            _cell_border(
                cell, "t", cb.get("t") or (tb.get("t") if i == 0 else None)
            )
            _cell_border(
                cell,
                "b",
                cb.get("b") or (tb.get("b") if i == len(rows) - 1 else None),
            )
            j += c["span"]
    return gf


def add_image(shapes, item: dict, base: Path):
    x, y, w, h = (item["rect"][k] for k in "xywh")
    src = item["src"]
    path = (
        Path(unquote(urlparse(src).path))
        if src.startswith("file:")
        else base / src
    )
    pic = shapes.add_picture(str(path), _emu(x), _emu(y), _emu(w), _emu(h))
    pic.name = "picture"
    if item["alt"]:
        pic._element.nvPicPr.cNvPr.set("descr", item["alt"])
    b = next(iter(item["borders"].values()), None)
    if b:
        pic.line.color.rgb = _rgb(b["c"])
        pic.line.width = _emu(b["w"])
    return pic


TITLE_NAMES = re.compile(r"^(h1|h2|when)$")


def build(slides: list[dict], out: Path, title: str, base: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    prs.core_properties.title = title
    prs.core_properties.author = "OCHA Centre for Humanitarian Data"
    blank = prs.slide_layouts[6]
    for s in slides:
        slide = prs.slides.add_slide(blank)
        titled = False
        for item in s["items"]:
            k = item["kind"]
            if k == "img":
                add_image(slide.shapes, item, base)
            elif k == "table":
                add_table(slide.shapes, item)
            elif k == "box":
                add_box(slide.shapes, item)
            else:
                shp = add_text(slide.shapes, item)
                if not titled and TITLE_NAMES.match(item.get("name", "")):
                    shp.name = "Title"
                    titled = True
    prs.save(out)


def export(
    index_html: Path, out: Path, title: str, chrome: str, dump: Path = None
):
    slides = extract(index_html, chrome)
    if dump:
        dump.write_text(json.dumps(slides, indent=1))
    build(slides, out, title, index_html.parent)
    return slides
